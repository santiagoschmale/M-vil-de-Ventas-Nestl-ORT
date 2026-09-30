"""
Importadores de las tres entradas del móvil. Reportan los problemas de calidad, no
los replican.

- Input 1: el Excel de Contraloría con el objetivo por SKU (SKU y kilos son
  esenciales; NNS opcional).
- Input 2: la tabla de totales por canal que pega el planner (texto separado por
  tabulaciones, punto y coma o comas).
- Base: el reparto del mes anterior, SKU × canal en kilos. Celda vacía = no aplica;
  0 = aplica con cero.
- Apertura: cómo se abrió cada celda debajo del canal el mes anterior (SKU, canal,
  entidad, kilos), en formato largo. Una entidad que no figura no aplica ahí.

Todo se detecta por nombre de encabezado, nunca por posición, y se cruza por
código de SKU, nunca por nombre. Los nombres aceptados están en formato.json.

Cada problema es un `Problema` con severidad "error" (el dato está mal y no se usa)
o "aviso" (raro pero válido). Si el archivo no se puede leer (falta una columna
esencial), se lanza ErrorDeEntrada.
"""

from __future__ import annotations

import csv
import io
import json
import re
import unicodedata
from dataclasses import dataclass, field
from decimal import ROUND_HALF_EVEN, Decimal
from pathlib import Path

from openpyxl import load_workbook

from src.domain.formato import a_texto
from src.domain.problema import Problema

FORMATO = json.loads(Path(__file__).with_name("formato.json").read_text(encoding="utf-8"))
KILOS = 3  # decimales
PLATA = 2


class ErrorDeEntrada(ValueError):
    """El archivo o la tabla no se puede leer: falta algo esencial."""


@dataclass
class ObjetivoSku:
    kilos: Decimal
    nns: Decimal | None
    descripcion: str
    categoria: str | None
    pais: str | None
    fila: int | None = None  # fila del Excel
    # Si el SKU está repetido: las otras filas, para que el planner elija cuál vale.
    alternativas: list[ObjetivoSku] = field(default_factory=list)


@dataclass
class TotalCanal:
    kilos: Decimal
    plata: Decimal | None


# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------

def normalizar(texto) -> str:
    """Minúsculas, sin tildes, espacios simples: para comparar encabezados."""
    texto = unicodedata.normalize("NFKD", str(texto))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", texto).strip().lower()


_MILES_AR = re.compile(r"-?\d{1,3}(\.\d{3})+(,\d+)?")
_COMA_DECIMAL = re.compile(r"-?\d+(,\d+)?")
_PUNTO_DECIMAL = re.compile(r"-?\d+(\.\d+)?")


def numero(texto: str) -> Decimal | None:
    """
    Número escrito a mano o pegado desde Excel -> Decimal, o None si no es número.

    Si tiene forma de separador de miles con punto (13.000, 1.234.567,89) se lee en
    formato argentino. Si no, el punto es decimal (1234.5). El caso ambiguo "1.234"
    se lee como mil doscientos treinta y cuatro: la tabla viene de un Excel en español.
    Se acepta el signo $ adelante. Un espacio en el medio no se adivina.
    """
    # Excel deja el signo $ y espacios duros al copiar celdas con formato moneda.
    t = str(texto).replace("\xa0", " ").strip()
    t = re.sub(r"^\$\s*", "", t).strip()
    if _MILES_AR.fullmatch(t):
        t = t.replace(".", "").replace(",", ".")
    elif _COMA_DECIMAL.fullmatch(t):
        t = t.replace(",", ".")
    elif not _PUNTO_DECIMAL.fullmatch(t):
        return None
    return Decimal(t)


def _celda(valor) -> tuple[Decimal | None, str | None]:
    """
    Celda de Excel -> (Decimal | None, problema | None). None = celda vacía.

    Excel guarda los números como float: `repr` da el decimal más corto que lo
    representa, que es lo que se ve en la celda. Es el único lugar donde entra un float.
    """
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None, None
    if isinstance(valor, bool):
        return None, f"valor no numérico ({valor!r})"
    if isinstance(valor, int):
        return Decimal(valor), None
    if isinstance(valor, float):
        return Decimal(repr(valor)), None
    n = numero(valor)
    return (n, None) if n is not None else (None, f"celda con error o texto ({valor!r})")


def _codigo(valor) -> str | None:
    if valor is None:
        return None
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = str(valor).strip()
    # Solo dígitos: 00123 y 123 son el mismo SKU (un archivo lo guarda como texto y otro como número).
    if texto.isdigit():
        texto = str(int(texto))
    return texto or None


def _columnas(encabezado, entrada: str) -> dict[str, int]:
    """Encabezado -> {campo: índice} con los alias de esa entrada en formato.json."""
    alias = {normalizar(a): campo for campo, lista in FORMATO[entrada].items() for a in lista}
    encontradas: dict[str, int] = {}
    for i, celda in enumerate(encabezado):
        if celda is not None:
            campo = alias.get(normalizar(celda))
            if campo and campo not in encontradas:
                encontradas[campo] = i
    return encontradas


def _al_paso(valor: Decimal, decimales: int, que: str, sku, problemas: list[Problema]) -> Decimal:
    paso = Decimal(1).scaleb(-decimales)
    redondeado = valor.quantize(paso, rounding=ROUND_HALF_EVEN)
    if redondeado != valor:
        problemas.append(Problema("aviso", f"{que} con más de {decimales} decimales: {a_texto(valor)} se redondea a "
                                           f"{a_texto(redondeado)}.", sku=sku))
    return redondeado


def _hojas_de_excel(origen) -> list[list[list]]:
    try:
        libro = load_workbook(origen, read_only=True, data_only=True)
    except Exception as e:  # openpyxl lanza de todo según qué esté roto
        raise ErrorDeEntrada(f"No se pudo abrir el archivo como Excel (.xlsx): {e}") from e
    try:
        return [[list(f) for f in hoja.iter_rows(values_only=True)] for hoja in libro.worksheets]
    finally:
        libro.close()


def _encabezado(hojas: list[list[list]], entrada: str, obligatorios: set[str], que: str):
    """La primera hoja (y fila, dentro de las primeras 30) que tenga las columnas obligatorias."""
    for filas in hojas:
        for n, fila in enumerate(filas[:30]):
            cols = _columnas(fila, entrada)
            if obligatorios <= cols.keys():
                return filas, n, cols
    nombres = ", ".join(sorted(obligatorios))
    raise ErrorDeEntrada(f"{que}: no se encontró un encabezado con las columnas {nombres} en ninguna hoja "
                         f"(nombres aceptados en formato.json).")


# ---------------------------------------------------------------------------
# Input 1: objetivo por SKU
# ---------------------------------------------------------------------------

def leer_input1(origen) -> tuple[dict[str, ObjetivoSku], list[Problema]]:
    filas, n, cols = _encabezado(_hojas_de_excel(origen), "input1", {"sku", "kilos"}, "Objetivo de Contraloría")
    problemas: list[Problema] = []
    if "nns" not in cols:
        problemas.append(Problema("aviso", "El objetivo de Contraloría no tiene columna de NNS: se reparten solo los kilos.",
                                  bloque="input1"))

    def campo(fila, nombre):
        i = cols.get(nombre)
        return fila[i] if i is not None and i < len(fila) else None

    objetivos: dict[str, ObjetivoSku] = {}
    todos = problemas
    for numero_fila, fila in enumerate(filas[n + 1:], start=n + 2):
        codigo = _codigo(campo(fila, "sku"))
        if codigo is None:
            continue
        repetida = codigo in objetivos
        # ponytail: los problemas de una fila repetida no se informan (nadie la eligió);
        # si se la elige, sus datos entran igual. Guardarlos si hace falta avisarlos después.
        problemas = [] if repetida else todos
        montos = {}
        for nombre, decimales, etiqueta in (("kilos", KILOS, "Kilos"), ("nns", PLATA, "NNS")):
            if nombre not in cols:
                montos[nombre] = None
                continue
            valor, error = _celda(campo(fila, nombre))
            if error:
                problemas.append(Problema("error", f"{etiqueta}: {error}. Se toma 0.", sku=codigo, bloque="input1"))
                valor = Decimal(0)
            elif valor is None:
                valor = Decimal(0)
            if valor < 0:
                problemas.append(Problema("error", f"{etiqueta} negativo ({a_texto(valor)}). Se toma 0.", sku=codigo,
                                          bloque="input1"))
                valor = Decimal(0)
            montos[nombre] = _al_paso(valor, decimales, etiqueta, codigo, problemas)

        pais = campo(fila, "pais")
        if pais and normalizar(pais) != normalizar(FORMATO["pais_local"]):
            problemas.append(Problema("aviso", f"Producto de otro país ({pais}) en el archivo de "
                                               f"{FORMATO['pais_local']}.", sku=codigo, bloque="input1"))
        if montos["kilos"] == 0 and not montos["nns"]:
            problemas.append(Problema("aviso", "SKU con objetivo en cero (obsoleto o estacional): no se reparte.",
                                      sku=codigo, bloque="input1"))
        elif montos["kilos"] == 0:
            problemas.append(Problema("error", f"SKU con NNS ({a_texto(montos['nns'])}) pero kilos en cero: los kilos son "
                                               f"esenciales, no se reparte ni en kilos ni en pesos.",
                                      sku=codigo, bloque="input1"))
        objetivo = ObjetivoSku(
            kilos=montos["kilos"],
            nns=montos["nns"],
            descripcion=str(campo(fila, "descripcion") or ""),
            categoria=campo(fila, "categoria"),
            pais=pais,
            fila=numero_fila,
        )
        if repetida:
            objetivos[codigo].alternativas.append(objetivo)
        else:
            objetivos[codigo] = objetivo
    for codigo, o in objetivos.items():
        if o.alternativas:
            filas_txt = _y([str(x.fila) for x in [o, *o.alternativas]])
            todos.append(Problema("error", f"Está repetido en el objetivo de Contraloría (filas {filas_txt}): se "
                                           f"usa la fila {o.fila} hasta que elijas cuál vale.",
                                  sku=codigo, bloque="input1", tipo="sku_repetido"))
    if not objetivos:
        raise ErrorDeEntrada("Objetivo de Contraloría: no hay ningún SKU.")
    return objetivos, todos


def _y(partes: list[str]) -> str:
    return partes[0] if len(partes) == 1 else f"{', '.join(partes[:-1])} y {partes[-1]}"


# ---------------------------------------------------------------------------
# Input 2: totales por canal
# ---------------------------------------------------------------------------

def leer_input2(origen) -> tuple[dict[str, TotalCanal], list[Problema]]:
    """
    Totales por canal: un Excel (archivo o ruta) o la tabla en texto (str), pegada
    o armada por la pantalla al editar. Los dos caminos validan igual.
    """
    if isinstance(origen, str):
        lineas = [linea for linea in origen.strip().splitlines() if linea.strip()]
        if not lineas:
            raise ErrorDeEntrada("Totales por canal: la tabla está vacía.")
        separador = "\t" if "\t" in lineas[0] else ";" if ";" in lineas[0] else ","
        hojas = [list(csv.reader(io.StringIO("\n".join(lineas)), delimiter=separador))]
    else:
        hojas = _hojas_de_excel(origen)
    filas, n, cols = _encabezado(hojas, "input2", {"canal", "kilos"}, "Totales por canal")
    ancho = len(filas[n])
    problemas: list[Problema] = []

    def _de(fila, nombre):
        return fila[cols[nombre]] if nombre in cols and cols[nombre] < len(fila) else None

    con_canal = [f for f in filas[n + 1:] if str(_de(f, "canal") or "").strip()]
    # Una columna de pesos entera vacía (p. ej. armados en la pantalla desde el mes anterior) es "sin pesos", no ceros.
    if "plata" in cols and all(not str(_de(f, "plata") or "").strip() for f in con_canal):
        del cols["plata"]
    if "plata" not in cols:
        problemas.append(Problema("aviso", "Los totales por canal no tienen pesos: se reparten solo los kilos.",
                                  bloque="input2"))

    canales: dict[str, TotalCanal] = {}
    vistos: set[str] = set()  # repetido con la misma regla que usa el cruce: sin tildes, mayúsculas ni espacios de más
    for fila in con_canal:
        canal = str(_de(fila, "canal")).strip()
        if normalizar(canal) in vistos:
            problemas.append(Problema("error", f"Canal {canal} repetido en los totales por canal: se usa la primera "
                                               f"fila.", bloque="input2"))
            continue
        if isinstance(origen, str) and len(fila) > ancho:
            # Típico de pegar separado por comas con montos 1.000,50: la coma decimal parte la celda.
            problemas.append(Problema("error", f"La fila de {canal} tiene más columnas que el encabezado: revisá "
                                               f"el separador (con coma decimal, separá con tabulaciones o punto y "
                                               f"coma). No se usa.", bloque="input2"))
            continue
        montos = {}
        for nombre, decimales in (("kilos", KILOS), ("plata", PLATA)):
            if nombre not in cols:
                montos[nombre] = None
                continue
            crudo = _de(fila, nombre)
            # Del texto llega "1.234,5"; de Excel, el número de la celda (1.234 es 1,234, no miles).
            valor, error = _celda(crudo)
            if valor is None and error is None:
                problemas.append(Problema("error", f"{nombre.capitalize()} de {canal} vacío. Se toma 0.",
                                          bloque="input2"))
                montos[nombre] = Decimal(0).quantize(Decimal(1).scaleb(-decimales))
                continue
            if valor is None or valor < 0:
                problemas.append(Problema("error", f"{nombre.capitalize()} de {canal} inválido ({crudo!r}). "
                                                   f"Se toma 0.", bloque="input2"))
                valor = Decimal(0)
            montos[nombre] = _al_paso(valor, decimales, f"{nombre.capitalize()} de {canal}", None, problemas)
        canales[canal] = TotalCanal(kilos=montos["kilos"], plata=montos["plata"])
        vistos.add(normalizar(canal))
    if not canales:
        raise ErrorDeEntrada("Totales por canal: no hay ningún canal con nombre.")
    return canales, problemas


# ---------------------------------------------------------------------------
# Base: reparto del mes anterior
# ---------------------------------------------------------------------------

def leer_base(origen) -> tuple[dict[tuple[str, str], Decimal], list[Problema]]:
    filas, n, cols = _encabezado(_hojas_de_excel(origen), "base", {"sku"}, "Base del mes anterior")
    meta = set(cols.values())
    canales = [(i, str(v).strip()) for i, v in enumerate(filas[n]) if v is not None and i not in meta]
    problemas: list[Problema] = []
    base: dict[tuple[str, str], Decimal] = {}
    vistos: set[str] = set()
    for fila in filas[n + 1:]:
        codigo = _codigo(fila[cols["sku"]]) if cols["sku"] < len(fila) else None
        if codigo is None:
            continue
        if codigo in vistos:
            problemas.append(Problema("error", "SKU repetido en la base: se usa la primera fila.", sku=codigo,
                                      bloque="base"))
            continue
        vistos.add(codigo)
        for i, canal in canales:
            valor, error = _celda(fila[i] if i < len(fila) else None)
            if error:
                problemas.append(Problema("error", f"{canal}: {error}. Se trata como no aplica.", sku=codigo,
                                          bloque="base"))
            elif valor is not None and valor < 0:
                problemas.append(Problema("error", f"Reparto negativo en {canal} ({a_texto(valor)}): es un error del "
                                                   f"archivo, no se usa.", sku=codigo, bloque="base"))
            elif valor is not None:
                base[(codigo, canal)] = valor
    if not base and not vistos:
        raise ErrorDeEntrada("Base del mes anterior: no hay ningún SKU.")
    return base, problemas



# ---------------------------------------------------------------------------
# Apertura debajo del canal
# ---------------------------------------------------------------------------

def leer_apertura(origen) -> tuple[dict[tuple[str, str], dict[str, Decimal]], list[Problema]]:
    """
    {(sku, canal): {entidad: peso}} con el reparto del mes anterior debajo de cada
    canal (distribuidores, vendedores). Supuesto A10: los vendedores cuelgan de la
    venta directa y de cada territorio.
    """
    filas, n, cols = _encabezado(_hojas_de_excel(origen), "apertura", {"sku", "canal", "entidad", "kilos"},
                                 "Apertura del mes anterior")
    problemas: list[Problema] = []
    apertura: dict[tuple[str, str], dict[str, Decimal]] = {}

    def campo(fila, nombre):
        i = cols[nombre]
        return fila[i] if i < len(fila) else None

    for fila in filas[n + 1:]:
        codigo = _codigo(campo(fila, "sku"))
        canal = str(campo(fila, "canal") or "").strip()
        entidad = str(campo(fila, "entidad") or "").strip()
        if codigo is None or not canal or not entidad:
            continue
        celda = apertura.setdefault((codigo, canal), {})
        if entidad in celda:
            problemas.append(Problema("error", f"{entidad} repetida en {canal}: se usa la primera fila.",
                                      sku=codigo, bloque="apertura"))
            continue
        valor, error = _celda(campo(fila, "kilos"))
        if error:
            problemas.append(Problema("error", f"{entidad} en {canal}: {error}. Se trata como no aplica.",
                                      sku=codigo, bloque="apertura"))
        elif valor is not None and valor < 0:
            problemas.append(Problema("error", f"Reparto negativo de {entidad} en {canal} ({a_texto(valor)}): es un error "
                                               f"del archivo, no se usa.", sku=codigo, bloque="apertura"))
        elif valor is not None:
            celda[entidad] = valor
    return {k: v for k, v in apertura.items() if v}, problemas
