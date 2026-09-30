"""
Recorrido de punta a punta, sin interfaz: input 1 + input 2 + base -> cruce en
kilos y en plata -> inconsistencias -> Excel.

    cd backend && .venv/bin/python -m src.movil.recorrido \\
        --input1 ../data/sample/input1_objetivo.xlsx \\
        --input2 ../data/sample/input2_canales.tsv \\
        --base ../data/sample/base_mes_anterior.xlsx \\
        --apagar 90020900 --salida movil.xlsx

Kilos y plata se cruzan por separado, cada uno contra su propio total ("cada unidad
se reparte y cuadra por separado"). Los dos usan la misma base del mes anterior: el
ajuste escala filas y columnas, así que el precio de cada canal lo absorbe solo.

Esta capa arma las entradas y presenta el resultado; las cuentas son del dominio.
"""

from __future__ import annotations

import argparse
import sys
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from src.domain.apertura import Nodo, recalcular
from src.domain.problema import Problema
from src.domain.cruce import Celda, Grupo, ResultadoCruce, cruzar
from src.importer.entradas import (
    KILOS,
    PLATA,
    ObjetivoSku,
    ErrorDeEntrada,
    TotalCanal,
    leer_apertura,
    leer_base,
    leer_input1,
    leer_input2,
    normalizar,
)


@dataclass(frozen=True)
class Regla:
    """
    Lo que define el planner: en `canal`, los SKUs de `categorias` suman a lo sumo
    (tope), al menos (mínimo) o exactamente (fijo) un % del total del canal. Un %
    para kilos y otro para NNS, independientes; None = en esa unidad no aplica.
    """
    id: str
    canal: str
    categorias: frozenset[str]
    limite: str  # tope | minimo | fijo
    kilos: Decimal | None
    nns: Decimal | None


@dataclass
class Recorrido:
    objetivos: dict[str, ObjetivoSku]
    canales: dict[str, TotalCanal]
    kilos: ResultadoCruce
    plata: ResultadoCruce | None  # None si algún input no trae plata
    problemas: list[Problema] = field(default_factory=list)
    # Debajo del canal: {(sku, canal): raíz con el valor del cruce y sus entidades}.
    aperturas: dict[Celda, Nodo] = field(default_factory=dict)
    entidades_apagadas: frozenset[str] = frozenset()


def armar(
    objetivos: dict[str, ObjetivoSku],
    canales: dict[str, TotalCanal],
    base: dict[Celda, Decimal],
    apagados: set[str] | frozenset[str] = frozenset(),
    fijas_kilos: dict[Celda, Decimal] | None = None,
    fijas_plata: dict[Celda, Decimal] | None = None,
    problemas: list[Problema] | None = None,
    apertura: dict[Celda, dict[str, Decimal]] | None = None,
    entidades_apagadas: set[str] | frozenset[str] = frozenset(),
    reglas: list[Regla] | tuple[Regla, ...] = (),
    canales_sku: dict[str, set[str] | frozenset[str]] | None = None,
    porcentajes: dict[tuple[str, str], Decimal] | None = None,
) -> Recorrido:
    """
    `apagados`: SKUs que el planner saca del mes (ON/OFF). No se reparten; si su
    objetivo estaba contado en el input 2, el cruce va a informar que los totales no
    coinciden, y es el planner quien ajusta el input 2.

    `apertura`: cómo se abre cada celda debajo del canal. Cada celda del cruce con
    apertura se reparte entre sus entidades con largest remainder, en kilos y en
    plata por separado. `entidades_apagadas` (distribuidores, vendedores) salen del
    reparto en todas las celdas y su parte va a los demás de la misma celda (A1).

    `reglas`: cada una se lleva a un grupo por unidad (kilos y NNS por separado),
    con los SKUs activos de sus categorías.

    `canales_sku`: dónde se vende cada SKU según el planner ("Dónde se vende"). Ver
    `_donde_se_vende`.

    `porcentajes`: base de cálculo, {(canal, entidad): % manual}. Esa entidad se lleva
    ese % de cada celda del canal que se abre; el resto va por histórico (supuesto B1).
    """
    problemas = list(problemas or [])
    apagados = frozenset(apagados)
    # Los canales se cruzan por nombre: sin importar tildes, mayúsculas ni espacios.
    por_nombre = {normalizar(c): c for c in canales}
    base = {(s, por_nombre.get(normalizar(c), c)): w for (s, c), w in base.items()}
    activos = {s: o for s, o in objetivos.items() if s not in apagados and o.kilos > 0}
    if canales_sku:
        base, apertura = _donde_se_vende(canales_sku, activos, canales, base, apertura or {}, por_nombre, problemas)

    for s in sorted({s for s, _ in base} - set(objetivos)):
        problemas.append(Problema("aviso", f"El SKU {s} está en la base del mes anterior pero no en el objetivo de Contraloría: "
                                           f"no se reparte.", sku=s, bloque="base"))
    for c in sorted({c for _, c in base} - set(canales)):
        problemas.append(Problema("aviso", f"El canal {c} está en la base del mes anterior pero no en los totales por canal: "
                                           f"no se reparte.", bloque="base"))

    grupos_kilos, grupos_plata = _grupos(reglas, activos, por_nombre, problemas)
    kilos = cruzar(
        filas={s: o.kilos for s, o in activos.items()},
        columnas={c: t.kilos for c, t in canales.items()},
        base=base,
        decimales=KILOS,
        fijas=fijas_kilos,
        grupos=grupos_kilos,
    )
    hay_plata = all(o.nns is not None for o in activos.values()) and all(t.plata is not None for t in canales.values())
    plata = None
    if hay_plata:
        plata = cruzar(
            filas={s: o.nns for s, o in activos.items() if o.nns > 0},
            columnas={c: t.plata for c, t in canales.items()},
            base=base,
            decimales=PLATA,
            fijas=fijas_plata,
            grupos=grupos_plata,
        )
    entidades_apagadas = frozenset(entidades_apagadas)
    aperturas = _abrir(kilos, plata, apertura or {}, entidades_apagadas, problemas, por_nombre, porcentajes or {})
    return Recorrido(objetivos=objetivos, canales=canales, kilos=kilos, plata=plata, problemas=problemas,
                     aperturas=aperturas, entidades_apagadas=entidades_apagadas)


def _donde_se_vende(canales_sku, activos, canales, base, apertura, por_nombre, problemas):
    """
    El planner dice en qué canales se vende un SKU: resuelve el SKU nuevo sin historia
    (A4) y restringe uno que ya tenía ("solo se vende en ..."). Devuelve base y apertura
    nuevas; no toca las de entrada.

    SUPUESTO (A4, a confirmar con el cliente): en un canal nuevo para el SKU se parte de
    objetivo del SKU × total del canal / total del mes (la tabla de independencia: un
    peso neutral, proporcional al canal), y el ajuste de siempre hace el resto. Debajo
    del canal, se abre como se abrió el canal entero el mes anterior (la suma de los
    pesos de cada entidad en ese canal); si no, esos kilos no llegarían a nadie.
    """
    base, apertura = dict(base), dict(apertura)
    total = sum(o.kilos for o in activos.values())
    for sku in sorted(canales_sku):
        if sku not in activos or not total:
            continue
        elegidos = set()
        for nombre in sorted(canales_sku[sku]):
            canal = por_nombre.get(normalizar(nombre))
            if canal is None:
                problemas.append(Problema("aviso", f"Dice que se vende en {nombre}, que no está en los totales por "
                                                   f"canal de este mes: ese canal no se usa.", sku=sku, bloque="canales"))
            else:
                elegidos.add(canal)
        if not elegidos:
            continue
        for celda in [k for k in base if k[0] == sku and k[1] not in elegidos]:
            del base[celda]  # "solo se vende en": los demás canales pasan a no aplica
        for canal in sorted(elegidos - {k for s, k in base if s == sku}):
            base[(sku, canal)] = activos[sku].kilos * canales[canal].kilos / total
            pesos: dict[str, Decimal] = {}
            for (x, k), entidades in apertura.items():
                if por_nombre.get(normalizar(k), k) == canal:
                    for e, w in entidades.items():
                        pesos[e] = pesos.get(e, Decimal(0)) + w
            if pesos:
                apertura[(sku, canal)] = pesos
    return base, apertura


def _grupos(reglas, activos, por_nombre, problemas) -> tuple[list[Grupo], list[Grupo]]:
    """
    Cada regla, en los SKUs activos de sus categorías. Canales y categorías se
    cruzan sin importar tildes ni mayúsculas. Una regla de un canal que no está en
    el input 2 no se aplica y se avisa; una categoría sin SKUs se avisa, pero la
    regla se aplica igual (un mínimo sobre nada es imposible y lo informa el cruce).
    """
    kilos, plata = [], []
    categorias = {normalizar(o.categoria or "") for o in activos.values()}
    for regla in sorted(reglas, key=lambda r: r.id):
        canal = por_nombre.get(normalizar(regla.canal))
        if canal is None:
            problemas.append(Problema("aviso", f"La regla {regla.id} es del canal {regla.canal}, que no está en el "
                                               f"los totales por canal de este mes: no se aplica.", bloque="reglas"))
            continue
        pedidas = {normalizar(c): c for c in regla.categorias}
        vacias = sorted(c for n, c in pedidas.items() if n not in categorias)
        if vacias:
            problemas.append(Problema("aviso", f"La regla {regla.id} incluye {', '.join(vacias)}, pero ningún SKU "
                                               f"activo del objetivo de Contraloría es de esa categoría.", bloque="reglas"))
        skus = frozenset(s for s, o in activos.items() if normalizar(o.categoria or "") in pedidas)
        for porcentaje, destino in ((regla.kilos, kilos), (regla.nns, plata)):
            if porcentaje is not None:
                destino.append(Grupo(regla.id, canal, skus, regla.limite, porcentaje))
    return kilos, plata


def _abrir(kilos, plata, apertura, apagadas, problemas, por_nombre, porcentajes) -> dict[Celda, Nodo]:
    """Una raíz por celda del cruce con apertura: el valor del cruce se reparte entre sus entidades."""
    if kilos.celdas is None:
        return {}
    aperturas = {}
    inactivas = {e: True for e in apagadas}
    for (sku, canal_base), pesos in sorted(apertura.items()):
        canal = por_nombre.get(normalizar(canal_base), canal_base)
        k = (sku, canal)
        valor_kilos = kilos.celdas.get(k, Decimal(0))
        valor_plata = Decimal(0) if plata is None or plata.celdas is None else plata.celdas.get(k, Decimal(0))
        if valor_kilos == 0 and valor_plata == 0:
            continue
        raiz = Nodo(entidad=f"{sku}|{canal}", nombre=canal, peso=Decimal(1),
                    hijos=[Nodo(entidad=e, nombre=e, peso=w, porcentaje=porcentajes.get((canal, e)))
                           for e, w in sorted(_con_porcentaje(pesos, canal, porcentajes).items())])
        raiz.valores["kilos"].monto = valor_kilos
        raiz.valores["nns"].monto = valor_plata
        recalcular(raiz, inactivas)
        if raiz.aviso["kilos"] or raiz.aviso["nns"]:
            problemas.append(Problema("error", por_que_no_abre(canal, raiz.aviso), sku=sku, bloque="apertura",
                                      canal=canal))
        aperturas[k] = raiz
    return aperturas


def _con_porcentaje(pesos: dict[str, Decimal], canal: str, porcentajes) -> dict[str, Decimal]:
    """
    Una entidad con % manual en el canal entra en todas sus celdas, aunque el mes anterior
    no haya vendido ese SKU (peso 0): su historia ya no la representa (supuesto B1).
    """
    return pesos | {e: Decimal(0) for (c, e) in porcentajes if c == canal and e not in pesos}


def por_que_no_abre(canal: str, avisos: dict[str, str | None]) -> str:
    """
    Los avisos del reparto de una celda ({"kilos": ..., "nns": ...}) dichos en términos
    del planner: qué pasó y qué queda. Si kilos y pesos fallan por lo mismo, un solo aviso.
    """
    def motivo(aviso: str) -> str:
        if "pesos son cero" in aviso:
            return "el mes anterior ninguno de sus distribuidores o vendedores vendió este SKU"
        if "No hay entidades" in aviso:
            return "todos sus distribuidores o vendedores están apagados"
        return aviso.rstrip(".")

    fallan = {("kilos" if u == "kilos" else "pesos"): motivo(a) for u, a in avisos.items() if a}
    if len(set(fallan.values())) == 1:
        medidas = " ni en ".join(fallan)
        return f"{canal} no se puede abrir en {medidas}: {next(iter(fallan.values()))}. Quedan en el canal, sin abrir."
    detalle = "; ".join(f"en {m}, {x}" for m, x in fallan.items())
    return f"{canal} no se puede abrir: {detalle}. Quedan en el canal, sin abrir."


def desde_archivos(input1, input2_texto: str, base_origen, **opciones) -> Recorrido:
    objetivos, p1 = leer_input1(input1)
    canales, p2 = leer_input2(input2_texto)
    base, p3 = leer_base(base_origen)
    try:
        apertura, p4 = leer_apertura(base_origen)
    except ErrorDeEntrada:
        apertura, p4 = {}, [Problema("aviso", "La base no tiene apertura debajo del canal: los canales cierran "
                                              "en el canal.", bloque="apertura")]
    return armar(objetivos, canales, base, problemas=p1 + p2 + p3 + p4, apertura=apertura, **opciones)


# ---------------------------------------------------------------------------
# Salida en Excel
# ---------------------------------------------------------------------------

def exportar(recorrido: Recorrido, destino) -> None:
    """
    Una hoja por unidad (SKU × canal, con los totales para verificar) y una hoja
    con los problemas de los archivos y las inconsistencias. Si el cruce no cerró,
    la hoja de esa unidad queda vacía y el motivo está en Problemas.
    """
    libro = Workbook()
    libro.remove(libro.active)
    canales = sorted(recorrido.canales)
    for nombre, cruce, total_sku, total_canal, formato in (
        ("Kilos", recorrido.kilos, lambda o: o.kilos, lambda t: t.kilos, "#,##0.000"),
        ("Plata", recorrido.plata, lambda o: o.nns, lambda t: t.plata, "#,##0.00"),
    ):
        hoja = libro.create_sheet(nombre)
        hoja.append(["Código SKU", "Descripción", "Objetivo"] + canales)
        for c in hoja[1]:
            c.font = Font(bold=True)
        if cruce is None or cruce.celdas is None:
            hoja.append(["Sin reparto: ver la hoja Problemas."])
            continue
        skus = sorted({s for s, _ in cruce.celdas})
        for s in skus:
            o = recorrido.objetivos[s]
            fila = [s, o.descripcion, total_sku(o)] + [cruce.celdas.get((s, c)) for c in canales]
            hoja.append(fila)
        hoja.append(["Total del reparto", None, sum((total_sku(recorrido.objetivos[s]) for s in skus), Decimal(0))]
                    + [sum((v for (_, x), v in cruce.celdas.items() if x == c), Decimal(0)) for c in canales])
        hoja.append(["Total pedido", None, None] + [total_canal(recorrido.canales[c]) for c in canales])
        for fila in hoja.iter_rows(min_row=2, min_col=3):
            for celda in fila:
                celda.number_format = formato

    hoja = libro.create_sheet("Apertura")
    hoja.append(["Código SKU", "Canal", "Entidad", "Kilos", "Plata", "Estado"])
    for c in hoja[1]:
        c.font = Font(bold=True)
    for (sku, canal), raiz in sorted(recorrido.aperturas.items()):
        for h in raiz.hijos:
            hoja.append([sku, canal, h.entidad, h.valores["kilos"].monto, h.valores["nns"].monto,
                         "activo" if h.activo else "apagado"])
    for fila in hoja.iter_rows(min_row=2, min_col=4, max_col=5):
        fila[0].number_format, fila[1].number_format = "#,##0.000", "#,##0.00"

    hoja = libro.create_sheet("Problemas")
    hoja.append(["Origen", "Severidad", "SKU", "Canal", "Detalle"])
    for c in hoja[1]:
        c.font = Font(bold=True)
    for unidad, cruce in (("kilos", recorrido.kilos), ("plata", recorrido.plata)):
        if cruce is None:
            continue
        for i in cruce.inconsistencias:
            hoja.append([f"cruce {unidad}", "error", ", ".join(i.skus), ", ".join(i.canales), i.mensaje])
        for a in cruce.avisos:
            hoja.append([f"cruce {unidad}", "aviso", None, None, a])
    for p in recorrido.problemas:
        hoja.append([p.bloque or "archivo", p.severidad, p.sku, p.canal, p.mensaje])
    libro.save(destino)


# ---------------------------------------------------------------------------
# Línea de comandos
# ---------------------------------------------------------------------------

def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Recorrido del móvil: inputs -> cruce -> Excel.")
    parser.add_argument("--input1", required=True, type=Path, help="Excel de Contraloría")
    parser.add_argument("--input2", required=True, type=Path, help="Tabla de totales por canal (texto)")
    parser.add_argument("--base", required=True, type=Path, help="Excel con el reparto del mes anterior")
    parser.add_argument("--apagar", action="append", default=[], help="Código de SKU a sacar del mes (repetible)")
    parser.add_argument("--salida", required=True, type=Path, help="Excel de salida")
    args = parser.parse_args(argv)

    r = desde_archivos(args.input1, args.input2.read_text(encoding="utf-8"), args.base, apagados=set(args.apagar))
    exportar(r, args.salida)
    for unidad, cruce in (("Kilos", r.kilos), ("Plata", r.plata)):
        if cruce is None:
            print(f"{unidad}: no se cruzó (falta la columna en algún input).")
        elif cruce.celdas is None:
            print(f"{unidad}: NO CIERRA.")
            for i in cruce.inconsistencias:
                print(f"  - {i.mensaje}")
        else:
            print(f"{unidad}: cierra exacto ({len(cruce.celdas)} celdas, {cruce.iteraciones} iteraciones).")
            for a in cruce.avisos:
                print(f"  aviso: {a}")
    errores = sum(1 for p in r.problemas if p.severidad == "error")
    print(f"Archivos: {errores} errores, {len(r.problemas) - errores} avisos. Detalle en {args.salida}.")
    cerro = r.kilos.celdas is not None and (r.plata is None or r.plata.celdas is not None)
    return 0 if cerro else 1


if __name__ == "__main__":
    sys.exit(main())
