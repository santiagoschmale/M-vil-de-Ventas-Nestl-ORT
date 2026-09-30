"""
API del móvil por HTTP: cargar entradas, ver inconsistencias, apagar, fijar y
exportar. Cada test arranca una app con una sesión vacía.
"""

from decimal import Decimal as D
from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import crear_app

MUESTRA = Path(__file__).parents[3] / "data" / "sample"
NUEVO = "90020900"


def _cliente():
    return TestClient(crear_app())


def _cargado(input2="input2_canales.tsv"):
    c = _cliente()
    with (MUESTRA / "input1_objetivo.xlsx").open("rb") as f:
        assert c.post("/api/movil/input1", files={"archivo": ("input1.xlsx", f)}).status_code == 200
    assert c.put("/api/movil/input2", json={"texto": (MUESTRA / input2).read_text(encoding="utf-8")}).status_code == 200
    with (MUESTRA / "base_mes_anterior.xlsx").open("rb") as f:
        assert c.post("/api/movil/base", files={"archivo": ("base.xlsx", f)}).status_code == 200
    return c


def _cerrado():
    c = _cargado("input2_sin_sku_nuevo.tsv")
    assert c.put(f"/api/movil/skus/{NUEVO}", json={"activo": False, "motivo": "SKU nuevo"}).status_code == 200
    return c


def test_sin_entradas_dice_que_falta_todo():
    r = _cliente().get("/api/movil").json()
    assert r["faltan"] == ["objetivo de Contraloría", "totales por canal", "mes anterior"]
    assert r["cierra"] == {"kilos": False, "plata": False}


def test_con_la_muestra_no_cierra_y_dice_por_que():
    r = _cargado().get("/api/movil").json()
    assert r["faltan"] == []
    assert r["cierra"] == {"kilos": False, "plata": False}
    (inc,) = r["inconsistencias"]["kilos"]
    assert inc["tipo"] == "sku_sin_canal" and inc["skus"] == [NUEVO]
    assert any("repetido" in p["mensaje"] for p in r["problemas"])
    assert r["entradas"]["input2"]["canales"] == 9


def test_los_montos_viajan_como_texto_con_los_decimales_de_su_unidad():
    c = _cerrado()
    r = c.get("/api/movil").json()
    assert r["cierra"] == {"kilos": True, "plata": True}
    assert isinstance(r["entradas"]["input1"]["kilos"], str)
    assert len(r["entradas"]["input1"]["kilos"].split(".")[1]) == 3
    assert len(r["entradas"]["input2"]["plata"].split(".")[1]) == 2
    cruce = c.get("/api/movil/cruce/kilos").json()
    montos = [celda["monto"] for s in cruce["skus"] for celda in s["celdas"].values()]
    assert montos and all(isinstance(m, str) and len(m.split(".")[1]) == 3 for m in montos)


def test_el_cruce_cierra_exacto_filas_y_columnas_por_http():
    cruce = _cerrado().get("/api/movil/cruce/kilos").json()
    for s in cruce["skus"]:
        if s["activo"] and s["celdas"]:
            assert sum(D(v["monto"]) for v in s["celdas"].values()) == D(s["objetivo"]), s["codigo"]
    for canal in cruce["canales"]:
        assert canal["repartido"] == canal["pedido"], canal["nombre"]


def test_fijar_y_desfijar_una_celda_con_motivo():
    c = _cerrado()
    cruce = c.get("/api/movil/cruce/kilos").json()
    sku = next(s for s in cruce["skus"] if s["celdas"] and any(D(v["monto"]) > 100 for v in s["celdas"].values()))
    canal = next(k for k, v in sku["celdas"].items() if D(v["monto"]) > 100)
    url = f"/api/movil/celdas/kilos/{sku['codigo']}/{canal}"

    assert c.put(url, json={"monto": "100", "motivo": ""}).status_code == 422  # sin motivo
    r = c.put(url, json={"monto": "100", "motivo": "acuerdo comercial"})
    assert r.status_code == 200
    celda = next(s for s in c.get("/api/movil/cruce/kilos").json()["skus"] if s["codigo"] == sku["codigo"])["celdas"][canal]
    assert celda == {"monto": "100.000", "fijada": True}
    historial = c.get("/api/movil").json()["historial"]
    assert historial[-1]["accion"] == "fijar" and historial[-1]["motivo"] == "acuerdo comercial"
    assert historial[-1]["autor"] == "planner-local"

    assert c.request("DELETE", url, json={"motivo": "se cayó"}).status_code == 200
    celda = next(s for s in c.get("/api/movil/cruce/kilos").json()["skus"] if s["codigo"] == sku["codigo"])["celdas"][canal]
    assert celda["fijada"] is False


def test_ediciones_imposibles_dan_422_con_mensaje():
    c = _cerrado()
    cruce = c.get("/api/movil/cruce/kilos").json()
    sku = next(s for s in cruce["skus"] if s["celdas"])
    canal = next(iter(sku["celdas"]))
    for unidad, monto in (("kilos", "-1"), ("kilos", "999999999"), ("kilos", "1.0001"), ("litros", "1")):
        r = c.put(f"/api/movil/celdas/{unidad}/{sku['codigo']}/{canal}", json={"monto": monto, "motivo": "x"})
        assert r.status_code == 422, (unidad, monto)
        assert r.json()["detail"]
    r = c.put(f"/api/movil/celdas/kilos/{sku['codigo']}/{canal}", json={"monto": 10, "motivo": "x"})
    assert r.status_code == 422  # un número JSON: los montos van como texto


def test_apagar_un_distribuidor_y_ver_la_apertura():
    c = _cerrado()
    cruce = c.get("/api/movil/cruce/kilos").json()
    sku = next(s for s in cruce["skus"] if "Distribuidores" in s["celdas"])
    url = f"/api/movil/apertura/{sku['codigo']}/Distribuidores"
    antes = c.get(url).json()
    assert antes["cuadra"] == {"kilos": True, "plata": True}
    entidad = antes["entidades"][0]["nombre"]

    assert c.put(f"/api/movil/entidades/{entidad}", json={"activo": False, "motivo": "convocatoria"}).status_code == 200
    despues = c.get(url).json()
    apagada = next(e for e in despues["entidades"] if e["nombre"] == entidad)
    assert apagada["activo"] is False and D(apagada["kilos"]) == 0
    assert sum(D(e["kilos"]) for e in despues["entidades"]) == D(despues["kilos"])


def test_exportar_devuelve_un_excel():
    r = _cerrado().get("/api/movil/exportar")
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert {"Kilos", "Plata", "Apertura", "Problemas"} <= set(load_workbook(BytesIO(r.content)).sheetnames)


def test_la_muestra_se_carga_de_un_solo_llamado():
    c = _cliente()
    assert c.post("/api/movil/muestra").status_code == 200
    assert c.get("/api/movil").json()["faltan"] == []


def test_un_archivo_que_no_es_excel_da_422_y_no_pisa_nada():
    c = _cargado()
    r = c.post("/api/movil/input1", files={"archivo": ("x.xlsx", b"no es excel")})
    assert r.status_code == 422
    assert c.get("/api/movil").json()["entradas"]["input1"]["archivo"] == "input1.xlsx"


def test_pedir_el_cruce_sin_entradas_da_409():
    assert _cliente().get("/api/movil/cruce/kilos").status_code == 409


def test_un_nombre_con_barra_llega_a_la_sesion_y_no_da_404():
    """Un canal o distribuidor real puede llamarse "Soluciones/Directa": el %2F no corta la ruta."""
    c = _cerrado()
    sku = next(s["codigo"] for s in c.get("/api/movil/cruce/kilos").json()["skus"] if s["celdas"])
    for url in ("/api/movil/entidades/Red%2FCuyo", f"/api/movil/celdas/kilos/{sku}/Soluciones%2FDirecta"):
        cuerpo = {"activo": False, "motivo": "x"} if "entidades" in url else {"monto": "1", "motivo": "x"}
        r = c.put(url, json=cuerpo)
        assert r.status_code == 422, (url, r.status_code)
        assert "/" in r.json()["detail"]
    assert c.get(f"/api/movil/apertura/{sku}/Soluciones%2FDirecta").json()["detail"].endswith("no se abre debajo del canal.")


def test_ningun_mensaje_para_el_planner_habla_de_input_1_o_input_2():
    """El planner conoce "objetivo de Contraloría", "totales por canal" y "mes anterior"."""
    import re

    c = _cliente()
    textos = list(c.get("/api/movil").json()["faltan"])
    c.post("/api/movil/muestra")
    c.put(f"/api/movil/skus/{NUEVO}", json={"activo": False, "motivo": "nuevo"})
    e = c.post("/api/movil/reglas", json={"canal": "Tucumán", "categorias": ["Té"], "limite": "tope", "kilos": "5",
                                           "nns": None, "motivo": "x"}).json()
    textos += [p["mensaje"] for p in e["problemas"]]
    textos += [i["mensaje"] for u in ("kilos", "plata") for i in e["inconsistencias"][u]]
    textos += [c.put("/api/movil/celdas/kilos/1/NoExiste", json={"monto": "1", "motivo": "x"}).json()["detail"]]
    textos += [c.put("/api/movil/skus/no-existe", json={"activo": False, "motivo": "x"}).json()["detail"]]
    textos += [c.put("/api/movil/input2", json={"texto": "esto no es una tabla"}).json()["detail"]]
    textos += [c.post("/api/movil/input1", files={"archivo": ("x.xlsx", b"no")}).json()["detail"]]
    assert len(textos) > 10
    hablan = [t for t in textos if re.search(r"input\s*[12]", t, re.I)]
    assert hablan == [], hablan


def test_los_totales_por_canal_se_suben_en_excel_y_el_estado_trae_cada_canal():
    c = _cliente()
    with (MUESTRA / "input2_canales.xlsx").open("rb") as f:
        r = c.post("/api/movil/input2", files={"archivo": ("totales.xlsx", f)})
    assert r.status_code == 200, r.json()
    totales = r.json()["entradas"]["input2"]
    assert totales["archivo"] == "totales.xlsx" and totales["canales"] == 9
    catering = next(t for t in totales["detalle"] if t["canal"] == "Catering")
    assert catering == {"canal": "Catering", "kilos": "62392.800", "plata": "493090298.40"}


def test_elegir_la_fila_de_un_sku_repetido_por_http():
    c = _cargado()
    (rep,) = c.get("/api/movil").json()["repetidos"]
    assert rep["sku"] == "90020001" and rep["elegida"] == 4
    assert [f["fila"] for f in rep["filas"]] == [4, 51]
    assert rep["filas"][1]["kilos"] == "1.000"
    r = c.put("/api/movil/skus/90020001/fila", json={"fila": 51, "motivo": "la otra era de otro producto"})
    assert r.status_code == 200 and r.json()["repetidos"][0]["elegida"] == 51
    assert c.put("/api/movil/skus/90020001/fila", json={"fila": 9, "motivo": "x"}).status_code == 422


def test_donde_se_vende_por_http():
    c = _cargado()
    r = c.put(f"/api/movil/skus/{NUEVO}/canales", json={"canales": ["Directa (BA)"], "motivo": "lanzamiento en BA"})
    assert r.status_code == 200, r.json()
    assert r.json()["cierra"]["kilos"] is True
    (d,) = r.json()["canales_sku"]
    assert d["sku"] == NUEVO and d["canales"] == ["Directa (BA)"] and d["motivo"] == "lanzamiento en BA"
    assert c.put(f"/api/movil/skus/{NUEVO}/canales", json={"canales": [], "motivo": "x"}).status_code == 422
    r = c.request("DELETE", f"/api/movil/skus/{NUEVO}/canales", json={"motivo": "se postergó"})
    assert r.status_code == 200 and r.json()["canales_sku"] == []


def test_deshacer_el_ultimo_cambio_una_sola_vez():
    c = _cargado()
    assert c.get("/api/movil").json()["deshacer"]["accion"] == "cargar_base"
    c.put(f"/api/movil/skus/{NUEVO}", json={"activo": False, "motivo": "SKU nuevo"})

    r = c.post("/api/movil/deshacer")
    assert r.status_code == 200
    assert r.json()["apagados"]["skus"] == [] and r.json()["deshacer"] is None
    assert r.json()["historial"][-1]["accion"] == "deshacer"
    assert c.post("/api/movil/deshacer").status_code == 422


def test_aprobar_y_reabrir():
    c = _cargado()
    assert c.post("/api/movil/aprobar").status_code == 422  # no cierra
    c = _cerrado()
    assert c.post("/api/movil/aprobar").status_code == 422  # faltan las etapas
    assert [e["etapa"] for e in c.get("/api/movil").json()["etapas"]] == ["canal", "apertura"]
    assert c.post("/api/movil/etapas/apertura").status_code == 422  # en orden
    c.post("/api/movil/etapas/canal")
    r = c.post("/api/movil/etapas/apertura").json()
    assert all(e["revisada"] for e in r["etapas"])
    r = c.post("/api/movil/aprobar").json()
    assert r["aprobado"]["accion"] == "aprobar"
    assert c.put(f"/api/movil/skus/{NUEVO}", json={"activo": True, "motivo": "m"}).status_code == 422
    r = c.post("/api/movil/reabrir", json={"motivo": "faltó un acuerdo"}).json()
    assert r["aprobado"] is None and r["historial"][-1]["motivo"] == "faltó un acuerdo"


def test_porcentaje_manual_de_una_entidad():
    c = _cerrado()
    cuerpo = {"canal": "Córdoba", "entidad": "Nicolás Paz", "porcentaje": "30", "motivo": "cartera nueva"}
    r = c.put("/api/movil/porcentajes", json=cuerpo)
    assert r.status_code == 200
    (p,) = r.json()["porcentajes"]
    assert (p["canal"], p["entidad"], p["porcentaje"], p["motivo"]) == ("Córdoba", "Nicolás Paz", "30", "cartera nueva")
    assert r.json()["porcentaje_asignado"] == {"Córdoba": "30"}
    sku = next(f["codigo"] for f in c.get("/api/movil/cruce/kilos").json()["skus"]
               if D(f["celdas"].get("Córdoba", {}).get("monto", "0")) > 0)
    a = c.get(f"/api/movil/apertura/{sku}/Córdoba").json()
    assert {e["nombre"]: e["porcentaje"] for e in a["entidades"]}["Nicolás Paz"] == "30"
    assert c.put("/api/movil/porcentajes", json={**cuerpo, "porcentaje": "150"}).status_code == 422
    r = c.put("/api/movil/porcentajes", json={**cuerpo, "porcentaje": None, "motivo": "vuelve"})
    assert r.json()["porcentajes"] == []


def test_alta_y_baja_de_una_entidad_nueva():
    c = _cerrado()
    cuerpo = {"canal": "Córdoba", "entidad": "Vendedora Nueva", "porcentaje": "20", "motivo": "entró en octubre"}
    r = c.post("/api/movil/entidades", json=cuerpo)
    assert r.status_code == 200
    assert [(e["canal"], e["entidad"]) for e in r.json()["entidades_nuevas"]] == [("Córdoba", "Vendedora Nueva")]
    sku = next(f["codigo"] for f in c.get("/api/movil/cruce/kilos").json()["skus"]
               if D(f["celdas"].get("Córdoba", {}).get("monto", "0")) > 0)
    nueva = next(e for e in c.get(f"/api/movil/apertura/{sku}/Córdoba").json()["entidades"] if e["nombre"] == "Vendedora Nueva")
    assert nueva["nueva"] and nueva["porcentaje"] == "20"
    assert c.post("/api/movil/entidades", json=cuerpo).status_code == 422  # ya está
    r = c.request("DELETE", "/api/movil/entidades", json={"canal": "Córdoba", "entidad": "Vendedora Nueva", "motivo": "error"})
    assert r.status_code == 200 and r.json()["entidades_nuevas"] == []


def test_el_estado_trae_los_kilos_por_canal_del_mes_anterior():
    """Para armar los totales por canal en la pantalla, arrancando del mes anterior."""
    c = _cliente()
    with (MUESTRA / "base_mes_anterior.xlsx").open("rb") as f:
        c.post("/api/movil/base", files={"archivo": ("base.xlsx", f)})
    canales = c.get("/api/movil").json()["entradas"]["base"]["canales"]
    assert [x["canal"] for x in canales] == sorted(x["canal"] for x in canales) and len(canales) > 1
    assert all(D(x["kilos"]) > 0 and D(x["kilos"]).as_tuple().exponent == -3 for x in canales)
