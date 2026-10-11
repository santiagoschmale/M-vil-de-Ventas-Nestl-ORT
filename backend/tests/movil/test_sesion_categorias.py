"""
Categorías guardadas en la herramienta (B8): si el objetivo del mes no trae categoría,
se toma de lo que el planner asignó. Sin categoría, un SKU no entra en ninguna regla.
"""

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook

from tests.movil.test_sesion import CUANDO, P, _falla, _sesion

DESPUES = datetime(2026, 10, 2, 11, 0)


def _sin_categorias():
    """El objetivo de la muestra, pero en el template propuesto: sin columna de categoría."""
    from openpyxl import load_workbook
    from tests.movil.test_sesion import MUESTRA
    origen = load_workbook(MUESTRA / "input1_objetivo.xlsx").active
    filas = list(origen.iter_rows(values_only=True))
    n = next(i for i, f in enumerate(filas) if f and "Categoría" in f)
    sin = [i for i, v in enumerate(filas[n]) if v != "Categoría"]
    wb = Workbook()
    for f in filas[n:]:
        wb.active.append([f[i] for i in sin])
    salida = BytesIO()
    wb.save(salida)
    salida.seek(0)
    return salida


def test_sin_categoria_en_el_objetivo_se_avisa_una_sola_vez():
    s = _sesion()
    s.cargar_input1(_sin_categorias(), "template.xlsx", P, CUANDO)
    assert s.sin_categoria()
    avisos = [p for p in s.recorrido.problemas if p.tipo == "sin_categoria"]
    assert len(avisos) == 1 and str(len(s.sin_categoria())) in avisos[0].mensaje


def test_asignar_categorias_completa_el_objetivo():
    s = _sesion()
    s.cargar_input1(_sin_categorias(), "template.xlsx", P, CUANDO)
    faltan = s.sin_categoria()
    s.asignar_categorias({sku: "Café" for sku in faltan}, P, DESPUES, "carga inicial")
    assert s.sin_categoria() == [] and not [p for p in s.recorrido.problemas if p.tipo == "sin_categoria"]
    assert all(s.objetivos_vigentes()[sku].categoria == "Café" for sku in faltan)
    assert s.historial[-1].accion == "asignar_categorias" and s.historial[-1].motivo == "carga inicial"


def test_la_categoria_del_archivo_manda_sobre_la_asignada():
    s = _sesion()  # la muestra trae categoría
    sku, o = next(iter(s.objetivos_vigentes().items()))
    s.asignar_categorias({sku: "Otra"}, P, DESPUES, "m")
    assert s.objetivos_vigentes()[sku].categoria == o.categoria


def test_se_rechaza_lo_que_no_es_valido():
    s = _sesion()
    assert _falla(s.asignar_categorias, {}, P, DESPUES, "m")
    assert _falla(s.asignar_categorias, {"90020001": "  "}, P, DESPUES, "m")
    assert _falla(s.asignar_categorias, {"90020001": "Café"}, P, DESPUES, "")


def test_los_codigos_a_mano_se_normalizan_como_en_el_archivo():
    s = _sesion()
    s.cargar_input1(_sin_categorias(), "template.xlsx", P, CUANDO)
    sku = s.sin_categoria()[0]
    s.asignar_categorias({f"00{sku}": "Café"}, P, DESPUES, "con ceros adelante")
    assert sku not in s.sin_categoria()
