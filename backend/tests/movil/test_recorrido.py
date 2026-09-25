"""
Recorrido de punta a punta sobre los datos ficticios de data/sample:
input 1 + input 2 + base -> cruce en kilos y en plata -> inconsistencias -> Excel.

Los valores esperados se calculan desde las entradas, sin pasar por el código que
se prueba.
"""

from decimal import Decimal as D
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from src.importer.entradas import leer_base, leer_input1, leer_input2
from src.movil.recorrido import armar, desde_archivos, exportar

MUESTRA = Path(__file__).parents[3] / "data" / "sample"
NUEVO = "90020900"  # SKU con objetivo y sin reparto previo (A4)


def _entradas(input2="input2_canales.tsv"):
    objetivos, _ = leer_input1(MUESTRA / "input1_objetivo.xlsx")
    canales, _ = leer_input2((MUESTRA / input2).read_text(encoding="utf-8"))
    base, _ = leer_base(MUESTRA / "base_mes_anterior.xlsx")
    return objetivos, canales, base


def _sin_el_nuevo(canales):
    """Lo que haría el planner: apaga el SKU nuevo y le saca a Directa lo que traía."""
    ajustados = dict(canales)
    directa = ajustados["Directa (BA)"]
    ajustados["Directa (BA)"] = type(directa)(kilos=directa.kilos - D("10"), plata=directa.plata - D("125000"))
    return ajustados


def test_con_la_muestra_tal_cual_el_sku_nuevo_frena_el_cruce():
    r = desde_archivos(
        MUESTRA / "input1_objetivo.xlsx",
        (MUESTRA / "input2_canales.tsv").read_text(encoding="utf-8"),
        MUESTRA / "base_mes_anterior.xlsx",
    )
    assert r.kilos.celdas is None
    assert [(i.tipo, i.skus) for i in r.kilos.inconsistencias] == [("sku_sin_canal", (NUEVO,))]
    assert [(i.tipo, i.skus) for i in r.plata.inconsistencias] == [("sku_sin_canal", (NUEVO,))]
    # Los problemas de los archivos llegan junto con el resultado.
    assert any("repetido" in p.mensaje for p in r.problemas)
    assert any("otro país" in p.mensaje for p in r.problemas)


def test_apagando_el_sku_nuevo_filas_y_columnas_cierran_exacto_en_kilos_y_plata():
    objetivos, canales, base = _entradas()
    r = armar(objetivos, _sin_el_nuevo(canales), base, apagados={NUEVO})

    assert r.kilos.inconsistencias == [] and r.plata.inconsistencias == []
    for unidad, cruce, decimales in (("kilos", r.kilos, 3), ("plata", r.plata, 2)):
        esperado_fila = {s: (o.kilos if unidad == "kilos" else o.nns) for s, o in objetivos.items() if s != NUEVO}
        for sku, total in esperado_fila.items():
            suma = sum((v for (s, _), v in cruce.celdas.items() if s == sku), D(0))
            assert suma == total, (unidad, sku, suma, total)
        for canal, t in _sin_el_nuevo(canales).items():
            esperado = t.kilos if unidad == "kilos" else t.plata
            suma = sum((v for (_, c), v in cruce.celdas.items() if c == canal), D(0))
            assert suma == esperado, (unidad, canal, suma, esperado)
        assert all(v.as_tuple().exponent == -decimales for v in cruce.celdas.values()), unidad


def test_los_skus_en_cero_y_los_apagados_no_se_reparten():
    objetivos, canales, base = _entradas()
    r = armar(objetivos, _sin_el_nuevo(canales), base, apagados={NUEVO})
    en_cero = {s for s, o in objetivos.items() if o.kilos == 0}
    assert len(en_cero) == 4
    repartidos = {s for s, _ in r.kilos.celdas}
    assert not (repartidos & (en_cero | {NUEVO}))


def test_el_ejemplo_del_cliente_informa_cuanto_le_falta_a_distribuidores():
    """
    Distribuidores pide más de lo que pueden darle los SKUs que se venden ahí. El
    faltante esperado se calcula aparte, por fuerza bruta (condición de Hall sobre
    los 512 conjuntos de canales): la suma de lo informado tiene que ser exactamente
    el faltante real, ni un gramo más ni uno menos.
    """
    from itertools import combinations

    objetivos, canales, base = _entradas("input2_distribuidores_imposible.tsv")
    canales = _sin_el_nuevo(canales)
    r = armar(objetivos, canales, base, apagados={NUEVO})

    activos = {s: o.kilos for s, o in objetivos.items() if s != NUEVO and o.kilos > 0}
    vende = {(s, c) for (s, c), w in base.items() if w > 0 and s in activos and c in canales}
    faltante = max(
        sum(canales[c].kilos for c in grupo) - sum(v for s, v in activos.items() if any((s, c) in vende for c in grupo))
        for n in range(1, len(canales) + 1)
        for grupo in combinations(sorted(canales), n)
    )
    assert faltante > 0
    assert r.kilos.celdas is None
    assert all(i.tipo == "canales_sin_volumen" for i in r.kilos.inconsistencias)
    assert sum(i.diferencia for i in r.kilos.inconsistencias) == faltante
    assert any("Distribuidores" in i.canales for i in r.kilos.inconsistencias)
    for i in r.kilos.inconsistencias:
        # Cada grupo informado: lo que piden sus canales menos lo que suman sus SKUs.
        pueden = sum(activos[s] for s in i.skus)
        assert i.diferencia == sum(canales[c].kilos for c in i.canales) - pueden
        assert i.diferencia.as_tuple().exponent == -3


def test_las_entidades_de_la_base_que_no_estan_en_los_inputs_se_avisan():
    objetivos, canales, base = _entradas()
    base = dict(base)
    base[("99999999", "Catering")] = D(10)  # SKU que ya no está en el input 1
    base[("90020001", "Canal viejo")] = D(10)  # canal que ya no está en el input 2
    r = armar(objetivos, _sin_el_nuevo(canales), base, apagados={NUEVO})
    avisos = [p.mensaje for p in r.problemas if p.severidad == "aviso"]
    assert any("99999999" in a for a in avisos)
    assert any("Canal viejo" in a for a in avisos)
    assert r.kilos.inconsistencias == []


def test_el_excel_de_salida_tiene_los_mismos_numeros_que_el_cruce():
    objetivos, canales, base = _entradas()
    r = armar(objetivos, _sin_el_nuevo(canales), base, apagados={NUEVO})
    destino = BytesIO()
    exportar(r, destino)
    destino.seek(0)
    libro = load_workbook(destino, data_only=True)

    assert {"Kilos", "Plata", "Problemas"} <= set(libro.sheetnames)
    hoja = libro["Kilos"]
    filas = list(hoja.iter_rows(values_only=True))
    encabezado = filas[0]
    col = {nombre: i for i, nombre in enumerate(encabezado)}
    comparadas = 0
    for fila in filas[1:]:
        sku = str(fila[col["Código SKU"]]) if fila[col["Código SKU"]] is not None else None
        if sku is None or sku.startswith("Total"):
            continue
        for canal in canales:
            esperado = r.kilos.celdas.get((sku, canal))
            celda = fila[col[canal]]
            if esperado is None:
                assert celda is None, (sku, canal)  # no aplica: celda vacía
            else:
                assert D(str(celda)) == esperado, (sku, canal, celda, esperado)
                comparadas += 1
    assert comparadas == len(r.kilos.celdas)


def test_con_inconsistencias_el_excel_igual_se_genera_con_la_lista_de_problemas():
    objetivos, canales, base = _entradas()
    r = armar(objetivos, canales, base)  # el nuevo encendido: no cierra
    destino = BytesIO()
    exportar(r, destino)
    destino.seek(0)
    problemas = list(load_workbook(destino)["Problemas"].iter_rows(values_only=True))
    assert any(NUEVO in " ".join(str(x) for x in fila if x) for fila in problemas)


def test_los_nombres_de_canal_de_la_base_cruzan_sin_importar_tildes_ni_mayusculas():
    from src.importer.entradas import ObjetivoSku, TotalCanal

    objetivos = {"1": ObjetivoSku(D("10"), D("100"), "", None, None)}
    canales = {"Córdoba": TotalCanal(D("10"), D("100"))}
    r = armar(objetivos, canales, {("1", "cordoba "): D(1)})
    assert r.kilos.celdas == {("1", "Córdoba"): D("10.000")}
    assert not any("cordoba" in p.mensaje for p in r.problemas)


# ---------------------------------------------------------------------------
# Apertura debajo del canal
# ---------------------------------------------------------------------------

def _con_apertura(apagados_entidades=frozenset()):
    from src.importer.entradas import leer_apertura

    objetivos, canales, base = _entradas()
    apertura, _ = leer_apertura(MUESTRA / "base_mes_anterior.xlsx")
    r = armar(objetivos, _sin_el_nuevo(canales), base, apagados={NUEVO}, apertura=apertura,
              entidades_apagadas=apagados_entidades)
    return r, apertura


def test_cada_celda_abierta_suma_exacto_o_queda_reportada():
    """
    Cada celda con apertura: o sus entidades suman exacto el valor del cruce (en kilos
    y en plata, con los decimales de cada unidad), o no cierra y está reportada con
    su motivo. Nunca una apertura que no suma sin que nadie lo diga.

    En la muestra hay celdas cuyas entidades tenían todas 0 el mes anterior (entidad
    sin histórico, A4): no hay base para repartir y tiene que quedar reportado.
    """
    r, apertura = _con_apertura()
    cierran = reportadas = 0
    for (sku, canal), raiz in r.aperturas.items():
        for unidad, cruce, decimales in (("kilos", r.kilos, 3), ("nns", r.plata, 2)):
            total = cruce.celdas.get((sku, canal), D(0))
            partes = [h.valores[unidad].monto for h in raiz.hijos]
            assert raiz.valores[unidad].monto == total
            assert all(p.as_tuple().exponent == -decimales for p in partes), (sku, canal, unidad, partes)
            if raiz.cuadra[unidad]:
                assert sum(partes, D(0)) == total, (sku, canal, unidad)
                cierran += 1
            else:
                assert raiz.aviso[unidad]
                assert any(p.sku == sku and canal in p.mensaje for p in r.problemas), (sku, canal)
                reportadas += 1
    assert cierran > 40 and reportadas > 0
    esperadas = {k for k in apertura if r.kilos.celdas.get(k, D(0)) > 0 or r.plata.celdas.get(k, D(0)) > 0}
    assert set(r.aperturas) == esperadas


def test_los_canales_sin_apertura_cierran_en_el_canal():
    r, _ = _con_apertura()
    assert not any(c in ("Catering", "Vending", "Mayoristas", "KAM Ingredientes", "KAM Sol") for _, c in r.aperturas)


def test_apagar_un_distribuidor_reparte_su_parte_entre_los_demas_en_cada_celda():
    antes, _ = _con_apertura()
    despues, _ = _con_apertura(apagados_entidades={"Red Cuyo"})
    tocadas = 0
    for k, raiz in despues.aperturas.items():
        nombres = {h.entidad: h for h in raiz.hijos}
        if "Red Cuyo" not in nombres:
            continue
        assert nombres["Red Cuyo"].valores["kilos"].monto == 0
        assert sum((h.valores["kilos"].monto for h in raiz.hijos), D(0)) == raiz.valores["kilos"].monto
        if antes.aperturas[k].hijos and any(
            h.entidad == "Red Cuyo" and h.valores["kilos"].monto > 0 for h in antes.aperturas[k].hijos
        ):
            tocadas += 1
    assert tocadas > 0
    # El cruce no cambia: apagar un distribuidor es debajo del canal.
    assert despues.kilos.celdas == antes.kilos.celdas


def test_la_apertura_sale_en_el_excel():
    r, _ = _con_apertura()
    destino = BytesIO()
    exportar(r, destino)
    destino.seek(0)
    filas = list(load_workbook(destino)["Apertura"].iter_rows(values_only=True))
    col = {n: i for i, n in enumerate(filas[0])}
    por_celda = {}
    for f in filas[1:]:
        clave = (str(f[col["Código SKU"]]), f[col["Canal"]])
        por_celda[clave] = por_celda.get(clave, D(0)) + D(str(f[col["Kilos"]]))
    # El Excel tiene exactamente lo que dio el reparto debajo de cada canal.
    assert set(por_celda) == set(r.aperturas)
    for k, total in por_celda.items():
        assert total == sum((h.valores["kilos"].monto for h in r.aperturas[k].hijos), D(0)), k


def test_por_que_no_abre_habla_en_terminos_del_planner():
    from src.movil.recorrido import por_que_no_abre

    m = por_que_no_abre("Rosario", "plata", "Todos los pesos son cero: no hay base para repartir.")
    assert m == ("Rosario no se puede abrir en pesos: el mes anterior ninguno de sus distribuidores o "
                 "vendedores vendió este SKU. Los pesos quedan en el canal, sin abrir.")
    assert "apagados" in por_que_no_abre("Directa (BA)", "kilos", "No hay entidades entre las cuales repartir.")
    assert "pesos son cero" not in m and "plata" not in m
