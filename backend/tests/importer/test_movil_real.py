"""
El móvil como lo arma hoy el planner (un Excel por bloques), leído como mes anterior.
El archivo de estos tests es inventado: tiene la forma del real, no sus datos.
"""

from decimal import Decimal as D
from io import BytesIO

from openpyxl import Workbook

from src.importer import movil_real
from src.importer.entradas import leer_apertura, leer_base, leer_input1, libro

KG3 = ["Kgs", "NNS", "NNS C/IIBB"]


def _bloques(hoja, titulos, filas, extra=("Totales", "Diferencia")):
    """Una tabla por bloques: títulos arriba, "SKU | Descripción | Kgs | NNS | NNS C/IIBB | | bloques…"."""
    arriba = [None] * 7
    encabezado = [None, "SKU", "Descripción", *KG3, None]
    for t in [*titulos, *extra]:
        arriba += [t, None, None]
        encabezado += KG3
    hoja.append(arriba)
    hoja.append(encabezado)
    for sku, desc, valores in filas:
        total = sum(valores)
        fila = [None, sku, desc, total, total * 10, total * 11, None]
        for v in [*valores, total, 0]:
            fila += [v, v * 10, v * 11]
        hoja.append(fila)


def _movil() -> BytesIO:
    wb = Workbook()
    mes = wb.active
    mes.title = "Móvil Mes"
    for _ in range(4):
        mes.append([])
    mes.append(["Categoria", "SKU", "Product Name", "Kgs", "NNS", "NNS C/IIBB"])
    mes.append(["Café", 1001, "Café soluble", 100.5, 1005.123, 1100])
    mes.append(["Vasos", 2001, "Vaso 8oz", 40, 400, 440])
    mes.append(["Café", 3001, "Discontinuado", 0, 0, 0])

    part = wb.create_sheet("Participaciones")
    # Después de Totales, los mismos canales otra vez: es el bloque de %.
    _bloques(part, ["Catering", "Vending"], [(1001, "Café soluble", [70, 30])], extra=("Totales", "Catering"))

    ing = wb.create_sheet("Ingredientes")
    _bloques(ing, ["Catering", "Vending"], [(1001, "Café soluble", [70.5, 0])])

    sol = wb.create_sheet("Soluciones")
    _bloques(sol, ["Call Center", "Distribuidores"], [(2001, "Vaso 8oz", [25, 15])])
    sol.append([])
    sol.append([None, "PARTICIPACIÓN ASESORES + KAM"])
    # Asesores: el nombre en cada columna (KG y NNS), como en el real.
    sol.append([None] * 7 + ["Asesor A-SAP", "Asesor A-SAP", "Asesor B-SAP", "Asesor B-SAP"])
    sol.append([None, "SKU", "Descripción", *KG3, None, "KG", "NNS", "KG", "NNS"])
    sol.append([None, 2001, "Vaso 8oz", 25, 250, 275, None, 20, 200, 5, 50])
    sol.append([None, None, None, None, None, "Total general", None, 20, 200, 5, 50])

    dis = wb.create_sheet("Distribuidores")
    _bloques(dis, ["Distribuidor X", "Distribuidor Y"], [(2001, "Vaso 8oz", [15, 0])])

    archivo = BytesIO()
    wb.save(archivo)
    archivo.seek(0)
    return archivo


def test_reconoce_el_movil_por_sus_hojas():
    assert movil_real.es_movil(libro(_movil()))
    otro = Workbook()
    otro.active.append(["SKU", "Catering"])
    archivo = BytesIO()
    otro.save(archivo)
    archivo.seek(0)
    assert not movil_real.es_movil(libro(archivo))


def test_el_mes_anterior_sale_de_las_hojas_de_reparto():
    base, problemas = leer_base(_movil())
    # Ingredientes y la primera tabla de Soluciones; un cero no aplica; Totales y Diferencia no son canales.
    assert base == {("1001", "Catering"): D("70.5"), ("2001", "Call Center"): D("25"),
                    ("2001", "Distribuidores"): D("15")}
    assert problemas == []


def test_la_apertura_son_los_distribuidores_y_los_asesores_del_call_center():
    apertura, problemas = leer_apertura(_movil())
    assert apertura == {
        ("2001", "Distribuidores"): {"Distribuidor X": D("15")},
        ("2001", "Call Center"): {"Asesor A-SAP": D("20"), "Asesor B-SAP": D("5")},
    }
    assert problemas == []


def test_el_objetivo_sale_de_movil_mes_con_su_categoria():
    objetivos, _ = leer_input1(_movil())
    assert set(objetivos) == {"1001", "2001", "3001"}
    assert objetivos["1001"].kilos == D("100.500") and objetivos["1001"].categoria == "Café"
    assert objetivos["1001"].nns == D("1005.12")
    assert objetivos["2001"].descripcion == "Vaso 8oz"


def test_en_participaciones_solo_cuentan_los_bloques_de_venta_no_los_de_porcentaje():
    hojas = libro(_movil())
    (tabla,) = movil_real.tablas(hojas["Participaciones"])
    assert tabla.titulos == ["Catering", "Vending"]
    assert tabla.filas["1001"]["Vending"][0] == D("30")
