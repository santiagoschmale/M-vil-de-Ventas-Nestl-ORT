"""
Tests de los importadores: input 1 (Excel de Contraloría), input 2 (tabla pegada)
y base (reparto del mes anterior). Contra los datos ficticios de data/sample y
contra casos armados a mano.
"""

from decimal import Decimal as D
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook

from src.importer.entradas import ErrorDeEntrada, leer_base, leer_input1, leer_input2, numero

MUESTRA = Path(__file__).parents[3] / "data" / "sample"


def _mensajes(problemas, severidad=None):
    return [p.mensaje for p in problemas if severidad in (None, p.severidad)]


def _excel(filas):
    wb = Workbook()
    for f in filas:
        wb.active.append(f)
    datos = BytesIO()
    wb.save(datos)
    datos.seek(0)
    return datos


def _falla(fn, *args):
    try:
        fn(*args)
        return False
    except ErrorDeEntrada:
        return True


# ---------------------------------------------------------------------------
# Números pegados
# ---------------------------------------------------------------------------

def test_numeros_en_formato_argentino_y_con_punto_decimal():
    assert numero("13.000") == D("13000")
    assert numero("1.234.567,89") == D("1234567.89")
    assert numero("1234,5") == D("1234.5")
    assert numero("1234.5") == D("1234.5")
    assert numero("1234.567") == D("1234.567")  # 4 dígitos antes del punto: no es separador de miles
    assert numero(" 500 ") == D("500")
    assert numero("-12,5") == D("-12.5")


def test_lo_que_no_es_numero_no_se_adivina():
    for texto in ("abc", "", "1,2,3", "1.23.4", "#N/A", "12 kg"):
        assert numero(texto) is None, texto


# ---------------------------------------------------------------------------
# Input 1
# ---------------------------------------------------------------------------

def test_input1_de_muestra():
    objetivo, problemas = leer_input1(MUESTRA / "input1_objetivo.xlsx")
    assert len(objetivo) == 47  # 45 con base + el nuevo + el de otro país; el repetido no suma
    assert all(isinstance(o.kilos, D) for o in objetivo.values())
    assert any("repetido" in m for m in _mensajes(problemas, "error"))
    assert any("otro país" in m for m in _mensajes(problemas, "aviso"))
    assert sum(1 for m in _mensajes(problemas, "aviso") if "en cero" in m) == 4


def test_input1_detecta_columnas_por_nombre_en_cualquier_orden():
    objetivo, _ = leer_input1(_excel([["KILOS", "Nns", "Código  SKU"], [100, 2000, 1]]))
    assert objetivo["1"].kilos == D("100") and objetivo["1"].nns == D("2000")


def test_input1_sin_nns_se_puede_leer_porque_lo_esencial_es_sku_y_kilos():
    objetivo, problemas = leer_input1(_excel([["SKU", "Kilos"], ["1", 10]]))
    assert objetivo["1"].nns is None
    assert any("NNS" in m for m in _mensajes(problemas, "aviso"))


def test_input1_sin_kilos_no_se_puede_leer():
    assert _falla(leer_input1, _excel([["SKU", "NNS"], ["1", 10]]))


def test_input1_valores_rotos_se_reportan():
    objetivo, problemas = leer_input1(_excel([["SKU", "Kilos"], ["1", -5], ["2", "#N/A"], ["3", 1.23456]]))
    errores = _mensajes(problemas, "error")
    assert any("negativo" in m for m in errores)
    assert any("#N/A" in m for m in errores)
    assert objetivo["3"].kilos == D("1.235")  # se redondea al gramo y se avisa
    assert any("decimales" in m for m in _mensajes(problemas, "aviso"))


# ---------------------------------------------------------------------------
# Input 2
# ---------------------------------------------------------------------------

def test_input2_de_muestra():
    canales, problemas = leer_input2((MUESTRA / "input2_canales.tsv").read_text(encoding="utf-8"))
    assert len(canales) == 9
    assert problemas == []
    assert all(c.plata is not None for c in canales.values())


def test_input2_acepta_tabulaciones_punto_y_coma_o_comas():
    tsv = "Canal\tKilos\tPlata\nMayoristas\t40.000\t1.000.000,50\n"
    pyc = "Canal;Kilos;Plata\nMayoristas;40.000;1.000.000,50\n"
    for texto in (tsv, pyc):
        canales, _ = leer_input2(texto)
        assert canales["Mayoristas"].kilos == D("40000")
        assert canales["Mayoristas"].plata == D("1000000.50")


def test_input2_canal_repetido_y_numero_roto():
    canales, problemas = leer_input2("Canal\tKilos\nA\t10\nA\t20\nB\tmucho\n")
    errores = _mensajes(problemas, "error")
    assert canales["A"].kilos == D("10")
    assert any("repetido" in m for m in errores)
    assert any("B" in m for m in errores)


def test_input2_sin_encabezado_reconocible_no_se_puede_leer():
    assert _falla(leer_input2, "Mayoristas\t40000\n")


# ---------------------------------------------------------------------------
# Base
# ---------------------------------------------------------------------------

def test_base_de_muestra():
    base, problemas = leer_base(MUESTRA / "base_mes_anterior.xlsx")
    # Cuenta independiente: celdas no vacías del Excel, leído directo.
    from openpyxl import load_workbook

    filas = list(load_workbook(MUESTRA / "base_mes_anterior.xlsx").active.iter_rows(values_only=True))
    no_vacias = sum(1 for f in filas[1:] for v in f[2:] if v is not None)
    assert len(base) == no_vacias > 100
    assert any(v == 0 for v in base.values())  # aplica con cero
    assert problemas == []


def test_base_vacio_es_no_aplica_y_cero_es_aplica_con_cero():
    base, _ = leer_base(_excel([["Código SKU", "Descripción", "X", "Y", "Z"], [1, "algo", 10, None, 0]]))
    assert base == {("1", "X"): D("10"), ("1", "Z"): D("0")}


def test_base_negativos_y_textos_se_reportan_y_no_se_usan():
    base, problemas = leer_base(_excel([["SKU", "X", "Y"], ["1", -3, "#N/A"], ["1", 5, 5]]))
    errores = _mensajes(problemas, "error")
    assert base == {}
    assert any("negativ" in m for m in errores)
    assert any("#N/A" in m for m in errores)
    assert any("repetido" in m for m in errores)


# ---------------------------------------------------------------------------
# Hallazgos del code review
# ---------------------------------------------------------------------------

def test_csv_con_comas_y_montos_con_coma_decimal_no_pierde_centavos_en_silencio():
    """Canal,Kilos,Plata con 1.000.000,50 sin comillas parte el monto: la fila se rechaza y se avisa."""
    canales, problemas = leer_input2("Canal,Kilos,Plata\nMayoristas,40.000,1.000.000,50\nVending,10,20\n")
    assert "Mayoristas" not in canales
    assert any("Mayoristas" in m and "separador" in m for m in _mensajes(problemas, "error"))
    assert canales["Vending"].plata == D("20.00")


def test_montos_pegados_con_signo_pesos_y_espacio_duro():
    assert numero("$ 1.000.000,50") == D("1000000.50")
    assert numero("1\xa0000,5") is None  # espacio como separador de miles: ambiguo, no se adivina
    assert numero("\xa040.000\xa0") == D("40000")
    assert numero("$-12,5") == D("-12.5")


def test_input2_kilos_vacios_es_error_no_cero_silencioso():
    canales, problemas = leer_input2("Canal\tKilos\nRosario\t\nVending\t10\n")
    assert any("Rosario" in m and "vac" in m for m in _mensajes(problemas, "error"))


def test_sku_con_nns_y_sin_kilos_se_reporta():
    objetivo, problemas = leer_input1(_excel([["SKU", "Kilos", "NNS"], ["1", 0, 500000]]))
    assert any("NNS" in m and "kilos" in m.lower() for m in _mensajes(problemas, "error"))


def test_el_encabezado_se_busca_en_todas_las_hojas():
    wb = Workbook()
    wb.active.title = "Portada"
    wb.active.append(["Móvil de mayo"])
    hoja = wb.create_sheet("Objetivo")
    hoja.append(["SKU", "Kilos"])
    hoja.append(["1", 10])
    datos = BytesIO()
    wb.save(datos)
    datos.seek(0)
    objetivo, _ = leer_input1(datos)
    assert objetivo["1"].kilos == D("10")


def test_codigos_con_ceros_a_la_izquierda_cruzan_con_los_numericos():
    objetivo, _ = leer_input1(_excel([["SKU", "Kilos"], ["00123", 10]]))
    base, _ = leer_base(_excel([["SKU", "X"], [123, 5]]))
    assert set(objetivo) == {s for s, _ in base}


# ---------------------------------------------------------------------------
# Apertura debajo del canal (hoja "Apertura anterior" de la base)
# ---------------------------------------------------------------------------

def test_apertura_de_muestra():
    from openpyxl import load_workbook

    from src.importer.entradas import leer_apertura

    apertura, problemas = leer_apertura(MUESTRA / "base_mes_anterior.xlsx")
    filas = list(load_workbook(MUESTRA / "base_mes_anterior.xlsx")["Apertura anterior"].iter_rows(values_only=True))
    assert sum(len(e) for e in apertura.values()) == len(filas) - 1  # una entrada por fila, cuenta independiente
    assert problemas == []
    assert {c for _, c in apertura} == {"Distribuidores", "Directa (BA)", "Córdoba", "Rosario"}


def test_apertura_negativa_repetida_o_con_texto_se_reporta():
    from src.importer.entradas import leer_apertura

    libro = _excel([
        ["SKU", "Canal", "Entidad", "Kilos"],
        ["1", "Distribuidores", "A", 10],
        ["1", "Distribuidores", "A", 20],
        ["1", "Distribuidores", "B", -5],
        ["1", "Distribuidores", "C", "#N/A"],
    ])
    apertura, problemas = leer_apertura(libro)
    assert apertura == {("1", "Distribuidores"): {"A": D("10")}}
    errores = _mensajes(problemas, "error")
    assert any("repetid" in m for m in errores)
    assert any("negativ" in m for m in errores)
    assert any("#N/A" in m for m in errores)


# ---------------------------------------------------------------------------
# Totales por canal desde Excel
# ---------------------------------------------------------------------------

def _excel(filas) -> BytesIO:
    wb = Workbook()
    for fila in filas:
        wb.active.append(fila)
    datos = BytesIO()
    wb.save(datos)
    datos.seek(0)
    return datos


def test_los_totales_por_canal_en_excel_dan_lo_mismo_que_la_tabla_pegada():
    pegada, _ = leer_input2((MUESTRA / "input2_canales.tsv").read_text(encoding="utf-8"))
    excel, problemas = leer_input2(MUESTRA / "input2_canales.xlsx")
    assert excel == pegada
    assert problemas == []
    # Mismos exponentes: 3 decimales en kilos y 2 en plata, venga de donde venga.
    assert all(t.kilos.as_tuple().exponent == -3 and t.plata.as_tuple().exponent == -2 for t in excel.values())


def test_en_excel_un_numero_con_tres_decimales_no_se_lee_como_miles():
    """En texto "1.234" es mil doscientos treinta y cuatro; en una celda numérica de Excel es 1,234."""
    canales, _ = leer_input2(_excel([["Canal", "Kilos", "Plata"], ["Catering", 1.234, 10.5]]))
    assert canales["Catering"].kilos == D("1.234") and canales["Catering"].plata == D("10.50")


def test_en_excel_los_problemas_se_reportan_igual():
    canales, problemas = leer_input2(_excel([["Canal", "Kilos", "Plata"], ["Catering", None, 10], ["Vending", -1, 5]]))
    mensajes = [p.mensaje for p in problemas]
    assert any("Catering" in m and "vac" in m for m in mensajes)
    assert any("Vending" in m for m in mensajes)
    assert canales["Catering"].kilos == D("0.000")


def test_un_sku_repetido_dice_en_que_filas_y_guarda_todas_para_elegir():
    objetivos, problemas = leer_input1(MUESTRA / "input1_objetivo.xlsx")
    (p,) = [p for p in problemas if p.tipo == "sku_repetido"]
    assert p.sku == "90020001" and "filas 4 y 51" in p.mensaje
    o = objetivos["90020001"]
    assert o.fila == 4 and o.kilos == D("7961.420")  # mientras nadie elija, vale la primera
    (otra,) = o.alternativas
    assert otra.fila == 51 and otra.kilos == D("1.000") and "repetido" in otra.descripcion
